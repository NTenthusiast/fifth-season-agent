#!/usr/bin/env python3
"""生成终审答辩 Q&A、逐字稿与演练清单。"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "决赛材料"
OUT.mkdir(exist_ok=True)
OUTPUT = OUT / "第五个季节_终审答辩_QA与逐字稿.docx"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    # OOXML 规定 shd 必须位于 vAlign/tcMar 等元素之前。
    inserted = False
    for later_name in ("tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"):
        later = tc_pr.find(qn(f"w:{later_name}"))
        if later is not None:
            tc_pr.insert(tc_pr.index(later), shd)
            inserted = True
            break
    if not inserted:
        tc_pr.append(shd)


def set_cell_text(cell, text, *, bold=False, color="182236", size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_qa(doc, number, question, short, expand, evidence, boundary=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(2.2)
    table.columns[1].width = Cm(13.8)
    left, right = table.rows[0].cells
    shade(left, "7257D5")
    shade(right, "EEE9FF")
    set_cell_text(left, f"Q{number:02d}", bold=True, color="FFFFFF", size=11)
    set_cell_text(right, question, bold=True, color="30255F", size=11)
    rows = [
        ("先答一句", short),
        ("展开回答", expand),
        ("证据位置", evidence),
    ]
    if boundary:
        rows.append(("边界提醒", boundary))
    for label, value in rows:
        cells = table.add_row().cells
        shade(cells[0], "F3F4F8")
        set_cell_text(cells[0], label, bold=True, color="60708A")
        set_cell_text(cells[1], value, color="182236")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None:
    zoom.set(qn("w:percent"), "100")
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "微软雅黑"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.25
for name, size, color in [("Title", 28, "30255F"), ("Heading 1", 18, "30255F"), ("Heading 2", 14, "4E3E9B")]:
    styles[name].font.name = "微软雅黑"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

# 封面
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(90)
r = p.add_run("第五季·续种")
r.bold = True; r.font.size = Pt(34); r.font.color.rgb = RGBColor(48, 37, 95)
r.font.name = "微软雅黑"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("终审答辩 Q&A 与 4 分钟逐字稿")
r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(114, 87, 213)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("2026 EL 智能应用开发与创新大赛 · AI 智能体创新专项组").italic = True
doc.add_paragraph()
box = doc.add_table(rows=4, cols=2)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (a, b) in enumerate([
    ("团队", "第五个季节"),
    ("答辩结构", "4 分钟 PPT + 3 分钟评委问答"),
    ("本版重点", "真实证据、可解释匹配、授权边界、接续准备度"),
    ("事实边界", "外部跨届试点待开展，不把演示数据表述为长期成效"),
]):
    shade(box.cell(i, 0), "EEE9FF")
    set_cell_text(box.cell(i, 0), a, bold=True, color="4E3E9B")
    set_cell_text(box.cell(i, 1), b)
doc.add_paragraph()
p = doc.add_paragraph("版本：2026-09-17 · 决赛强化版")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

add_heading(doc, "一、上台前只记住这五句话", 1)
add_bullets(doc, [
    "我们解决的是校园项目因毕业、换届、时间不足而中断后，经验与责任链一起消失的问题。",
    "知颜管理的是传承生命周期，不是替人做项目：建档—执行—沉淀—休眠—接棒。",
    "AI 推荐不自动放行；匹配显示证据、缺口与不确定性，入组和接棒仍由人确认。",
    "交接包强制包含授权范围、资源有效期和原团队权益；接续准备度未满 100 分就阻断正式唤醒。",
    "当前证明流程可运行，长期效果必须通过真实社团的跨周期试点验证，我们不会用演示数据替代。",
])

add_heading(doc, "二、4 分钟逐字稿", 1)
script_rows = [
    ("00:00—00:15", "封面", "各位评委老师好，我们是“第五个季节”。我们关注的不是如何再做一个校园项目，而是如何让已经投入过时间和热爱的项目，不因毕业和换届而归零。"),
    ("00:15—00:45", "问题", "校园项目常常不是做失败了，而是负责人毕业、社团换届或阶段性时间不足。资料还在网盘里，但经验、授权边界和下一步行动一起消失。校园需要的不是墓碑式归档，而是可被重新启动的传承机制。"),
    ("00:45—01:20", "方案", "我们把项目生命周期设计成萌芽、执行、沉淀、休眠和第五季接棒。知颜通过自然语言完成建档、组队、动态记录、交接包、休眠和唤醒。每次接棒留下交接内容、匹配理由和后续动态三类证据。"),
    ("01:20—02:10", "核心演示", "用户提出想接棒休眠项目，知颜先读取完整档案，再解释能力命中和缺口；随后检查交接包五件套、过程动态、授权范围、资源有效期和原团队权益。只有准备度满分且本人确认，系统才变更负责人并生成传承故事。AI 推荐不等于 AI 替人决定。"),
    ("02:10—02:50", "技术创新", "匹配采用四维权重：技能 40%、兴趣议题 35%、传承准备度 15%、时间投入 10%。档案没有时间投入就明确标为未知。接续治理采用 100 分制，缺少任一治理项都阻断唤醒。"),
    ("02:50—03:35", "验证边界", "我们接受评审对数据规模和长期效果的质疑。目前能证明的是流程可运行、数据可留痕，不能把演示库说成真实跨届成效。下一步首轮招募三到五个真实社团项目，连续八周跟踪资料可用率、三十天活跃率和首个里程碑；样本少于十，只作案例结论。"),
    ("03:35—04:00", "结尾", "我们不承诺每一颗种子都能重生，但希望每一次交接都比从零开始更可靠：建档可追溯，匹配可解释，授权有边界，效果能复盘。第五季，让未完成的美好拥有下一位同行者。谢谢各位老师。"),
]
t = doc.add_table(rows=1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, head in enumerate(["时间", "页面", "逐字稿"]):
    shade(t.cell(0, i), "7257D5"); set_cell_text(t.cell(0, i), head, bold=True, color="FFFFFF")
for tm, page, text in script_rows:
    row = t.add_row().cells
    set_cell_text(row[0], tm, bold=True, color="4E3E9B")
    set_cell_text(row[1], page, bold=True)
    set_cell_text(row[2], text)
doc.add_paragraph("演练标准：正常语速 220—240 字/分钟；必须在 3 分 45 秒前进入结尾页，为现场停顿留出余量。")
doc.add_page_break()

add_heading(doc, "三、高概率评委问答", 1)
qas = [
    ("你们和普通网盘、知识库有什么区别？", "我们管理的是接续决策，不只是文件存储。", "网盘回答“文件在哪里”；第五季还回答“项目为什么暂停、资料是否可用、谁适合接、授权边界是什么、接棒后有没有继续推进”。", "PPT 第3—5页；src/tools/project_tools.py", None),
    ("真正的 AI 智能体体现在哪里？", "AI 负责跨工具理解、编排与解释，但关键决策保留给人。", "知颜将自然语言意图路由到 26 个工具，跨用户档案、项目状态、交接包与匹配记录生成行动链；同时受到状态机、权限和准备度校验约束。", "src/agents/agent.py；src/tools/__init__.py；config/agent_llm_config.json", "不要把“使用大模型”本身当创新点，重点讲受约束的工具编排。"),
    ("你们的匹配分会不会是拍脑袋？", "权重公开、证据可见、未知项不补分。", "技能 40%、兴趣议题 35%、传承准备度 15%、时间投入 10%；结果返回命中项、能力缺口、不确定性和决策边界。权重会在真实试点后根据失败案例修订。", "PPT 第5页；score_project_match；scripts/test_explainable_match.py", "不要说“算法准确率很高”，当前没有足够真实标签验证准确率。"),
    ("为什么工具这么多，不会调错吗？", "工具数量不是卖点，清晰路由和前置条件才是。", "系统提示词有意图—工具路由表；相近工具在文档中写明不适用场景；敏感动作由代码层再次校验，不能只靠模型自觉。", "config/agent_llm_config.json；各工具 docstring；src/web_server.py", None),
    ("你们说已经有唤醒记录，这能证明项目有效吗？", "不能，只能证明流程与留痕机制可运行。", "长期有效必须看接棒后 30 天活跃、首个里程碑和资料可用率。首页验证中心刻意把运行证据和跨周期试点分开。", "PPT 第6页；/api/evidence；docs/真实试点方案.md", "严禁把演示种子或运行记录说成真实长期成效。"),
    ("线上数据规模这么小，结论可靠吗？", "我们不从小样本推导普遍结论。", "首轮试点计划 3—5 个项目，先验证流程和指标可操作性；样本少于 10 时只报告案例、失败原因与失访，不报告普遍成功率。", "docs/真实试点方案.md 第2、4节", None),
    ("资料质量怎么保证？", "机器检查完整性，人确认真实性与可用性。", "系统要求五件套非空并保留过程动态；接棒者在试点中逐项评价可直接使用、需补充或不可使用。AI 只能整理已有事实，禁止编造。", "assess_transfer_readiness；docs/真实试点方案.md 第3节", None),
    ("如果资料有版权或隐私问题怎么办？", "默认不自动转移敏感资源，授权范围和有效期是必填项。", "交接包必须声明授权范围、资源有效期与原团队权益；账号密码、个人联系方式、未脱敏访谈资料不得自动进入交接包。", "PPT 第5页；save_handover_package；试点方案第5节", None),
    ("接棒人拿到项目后改得面目全非怎么办？", "交接不是版权转让，重大改编按授权约定再次确认。", "原团队保留署名和成果归属；接棒者只获得约定用途内权限，系统保留匹配理由和接棒记录，争议发生时暂停受限资料访问并人工处理。", "docs/真实试点方案.md 第5节", None),
    ("AI 推荐错人造成损失，责任是谁？", "AI 只做解释性推荐，不做最终授权。", "入组须发起人审批，接棒须本人确认；系统展示能力缺口和未知项，敏感操作由负责人权限和准备度校验共同把关。", "PPT 第4—5页；search_projects / awaken_project", None),
    ("为什么叫第五季？", "四季是项目生命周期，第五季是跨越暂停之后的重生阶段。", "春萌芽、夏生长、秋沉淀、冬蛰伏；第五季不是自然季节，而是让既有投入获得第二次机会的产品隐喻。", "PPT 第3页；首页品牌叙事", None),
    ("休眠是不是在美化失败？", "不是。休眠是可恢复状态，但失败原因仍需真实保留。", "系统记录休眠原因、踩坑和遗留问题；试点要求失败案例不删除，用于修订交接模板和匹配权重。", "advance_project_stage；docs/真实试点方案.md 第4节", None),
    ("如何防止演示数据污染真实数据？", "演示种子只用于流程展示，不计入真实试点样本。", "验证计划明确区分演示库与试点项目；所有成功率必须同时报告样本量与失访数。", "docs/真实试点方案.md 第2、4节", None),
    ("如果数据库或模型挂了怎么办？", "界面不会用缓存假数据冒充实时结果，关键业务留痕在数据库。", "验证中心查询失败时显示暂不可用；模型与数据库有错误兜底和连接重试，但断网状态下不承诺完整智能体能力。", "src/web_server.py；src/storage/database；改动报告.md", "不要声称完全离线可用。"),
    ("项目的最大技术风险是什么？", "长期效果指标和权限治理的真实运行，而不是页面功能。", "技术链路已能运行，下一阶段最大风险是样本获取、社团配合和 8 周持续跟踪；因此我们把试点设计、失败保留和人工复核放在优先级最高的位置。", "PPT 第6页；试点方案", None),
    ("未来如何规模化？", "先做校内轻量试点，再标准化交接模板与指标，不急于扩校。", "先验证 3—5 个项目，形成领域模板、授权分级和复盘机制；只有口径稳定后再扩展到学院、社团联合会或多校合作。", "docs/真实试点方案.md", None),
    ("商业价值在哪里？", "短期价值是降低组织换届损耗，长期可作为高校项目资产与学生成长基础设施。", "可服务社团换届、公益项目、课程实践和校园创新；商业化应建立在学校治理与隐私合规之上，当前阶段优先验证公共价值。", "项目文档；README", "不要在没有定价和采购验证时声称成熟商业模式。"),
    ("你们相比其他同类作品的创新点是什么？", "把“归档”改造成可唤醒的状态机，并把匹配、治理和验证放进同一闭环。", "单点推荐、知识库或项目管理工具各自只覆盖一段；第五季把沉淀、休眠、可解释匹配、授权校验和接棒后追踪连成闭环。", "PPT 第3—6页", None),
    ("如果只能保留一个功能，会保留什么？", "保留“有治理边界的交接包 + 接续准备度”。", "没有高质量且可授权的资料，再好的匹配也只是把新人送进一个信息黑洞；因此交接质量是整个系统的地基。", "save_handover_package；assess_transfer_readiness", None),
    ("决赛后第一件事是什么？", "锁定真实试点伙伴并完成第 0 周授权与基线记录。", "先联系 3—5 个换届中的社团或项目组，签署一页式信息分级确认单，记录传统交接耗时和遗漏项，再进入系统交接。", "docs/真实试点方案.md 第6节", None),
]
for idx, qa in enumerate(qas, 1):
    add_qa(doc, idx, *qa)

doc.add_page_break()
add_heading(doc, "四、评审意见逐条回应矩阵", 1)
matrix = [
    ("数据规模小、长期效果待验证", "新增验证中心；运行证据与长期成效分开；制定 8 周真实试点", "不宣称已验证长期有效"),
    ("依赖学生建档与接棒意愿", "匹配显示命中证据、能力缺口和未知项；接棒需本人确认", "不把推荐分当自动决策"),
    ("工具编排复杂", "意图路由表 + 工具不适用场景 + 代码前置校验", "工具数不是主要创新点"),
    ("交接资料可用性", "五件套完整性 + 过程动态 + 试点中的逐项可用性评价", "机器校验不替代人工真实性核验"),
    ("授权、有效期、原团队权益", "三项必填治理声明；准备度未满分阻断唤醒", "交接不等于版权转让"),
    ("实际匹配、推荐如何有效", "四维公开评分 + 证据/缺口/不确定性 + 保存匹配日志", "当前不声称有高准确率"),
    ("使用频次与效果验证", "真实试点跟踪资料可用率、30天活跃率、首个里程碑", "所有比例同时报告样本量和失访"),
]
t = doc.add_table(rows=1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["评审关切", "已落地改进", "答辩边界"]):
    shade(t.cell(0, i), "7257D5"); set_cell_text(t.cell(0, i), h, bold=True, color="FFFFFF")
for row_data in matrix:
    cells = t.add_row().cells
    for i, value in enumerate(row_data):
        set_cell_text(cells[i], value, bold=(i == 0))
        if i == 2: shade(cells[i], "FFF2E8")

add_heading(doc, "五、现场演练清单", 1)
add_bullets(doc, [
    "PPT 共 7 页；第 4 页演示闭环最多 50 秒，第 6 页验证边界必须讲到。",
    "准备一个已完整建档的测试用户、一个休眠项目、一个新版治理交接包；不要现场临时造数据。",
    "演示前清空无关浏览器标签，关闭消息通知；将页面和 PPT 都预先打开。",
    "若模型响应超过 8 秒，直接切回 PPT，用第 4 页四步流程继续讲，不在台上排错。",
    "若数据库不可用，明确说明“系统按规则不展示模拟值”，展示本地截图和代码验证结果。",
    "问答时先给结论，再给机制，最后给证据位置；每题控制在 20—30 秒。",
    "任何涉及效果、准确率、社团数量的问题，都先说明样本边界，绝不临场放大数字。",
])

add_heading(doc, "六、证据索引", 1)
for label, path_text in [
    ("可解释匹配", "src/tools/project_tools.py → score_project_match"),
    ("接续准备度", "src/tools/project_tools.py → assess_transfer_readiness"),
    ("交接治理强校验", "src/tools/project_tools.py → save_handover_package"),
    ("正式唤醒阻断", "src/tools/project_tools.py → awaken_project"),
    ("实时验证中心", "src/web_server.py → /api/evidence"),
    ("真实试点口径", "docs/真实试点方案.md"),
    ("无数据库测试", "scripts/test_explainable_match.py；scripts/test_transfer_readiness.py"),
]:
    p = doc.add_paragraph()
    p.add_run(label + "：").bold = True
    p.add_run(path_text)

# 页眉页脚
for table in doc.tables:
    widths = [2.2, 13.8] if len(table.columns) == 2 else [3.0, 3.0, 10.0]
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)

for section in doc.sections:
    header = section.header.paragraphs[0]
    header.text = "第五季·续种｜终审答辩内部材料"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(96, 112, 138)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("2026 EL AI 智能体创新专项组 · 第五个季节")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(96, 112, 138)

doc.save(OUTPUT)
print(OUTPUT)
